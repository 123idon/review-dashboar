"""백년화편 네이버 스마트스토어 후기 수집기 (회사 PC용, 판매자센터 세션 유지)

구조
  1) --login : 브라우저를 띄워 사람이 직접 판매자센터에 로그인 → 세션이 로컬 프로필에 저장됨
  2) 기본 실행 : 저장된 세션으로 판매자센터 "리뷰관리" 화면을 읽어 후기 수집 → 대시보드 서버에 업로드
     - 화면 읽기가 실패하면 "엑셀 다운로드" 폴백으로 같은 데이터를 확보
  3) 세션 만료 시 : 서버에 쿠키만료 알림(대시보드 빨간 배너) 후 종료 → --login 한 번 더 실행

설치 (최초 1회, PowerShell)
  pip install playwright openpyxl
  python -m playwright install chromium

로그인 후 서버로 세션 이전 (권장 — 이후 PC가 꺼져 있어도 서버가 매일 수집)
  set NAVER_CONNECT_TOKEN=<토큰>  →  python naver_seller_local.py --login --push

로그인만 (PC에서 직접 수집할 때)
  python naver_seller_local.py --login

매일 실행 (작업 스케줄러)
  python naver_seller_local.py

프록시·지문 위조·캡차 우회 같은 것은 일절 사용하지 않는다. 본인 계정으로 본인 스토어 화면을 읽을 뿐이다.
"""
import argparse
import asyncio
import json
import os
import re
import sys
import time
from datetime import datetime, timedelta, date
from pathlib import Path
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
from uuid import uuid4

BASE_URL = os.environ.get("REVIEW_DASHBOARD_URL", "https://web-production-ce7ca1.up.railway.app").rstrip("/")
STORE_NAME = "백년화편"
SELLER_HOME = "https://sell.smartstore.naver.com/"
REVIEW_URL = "https://sell.smartstore.naver.com/#/review/search"

APP_DIR = Path(os.environ.get("LOCALAPPDATA") or Path.home()) / "review_dashboard"
PROFILE_DIR = APP_DIR / "naver_profile"     # 브라우저 프로필(쿠키·로그인 유지) — 절대 공유하지 말 것
DOWNLOAD_DIR = APP_DIR / "downloads"
LOG_PATH = APP_DIR / "naver_seller_local.log"
LAST_RESULT = APP_DIR / "last_result.json"

ROW_JS = r"""() => Array.from(document.querySelectorAll('[role="row"][row-index]')).map(e=>{
 const cell=k=>e.querySelector('[col-id="'+k+'"]');const txt=k=>cell(k)?.textContent.trim()||'';
 const id=cell('reviewContent')?.querySelector('a')?.getAttribute('ng-click')?.match(/openReviewDetailModal\((\d+)/)?.[1];
 return id?{row_index:Number(e.getAttribute('row-index')),review_no:id,product_no:txt('productNo'),product:txt('productName'),score:Number(txt('reviewScore')),
 content:txt('reviewContent'),author:txt('writerId'),date:txt('createDate').slice(0,10).replaceAll('.','-'),
 review_type:txt('reviewType'),platform:'naver',title:''}:null}).filter(Boolean)"""


# ───────────────────────── 공통 유틸 ─────────────────────────
def log(msg):
    line = f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {msg}"
    print(line, flush=True)
    try:
        APP_DIR.mkdir(parents=True, exist_ok=True)
        with LOG_PATH.open("a", encoding="utf-8") as f:
            f.write(line + "\n")
    except OSError:
        pass


class SessionExpired(Exception):
    pass


def api(path, body=None, timeout=90):
    payload = json.dumps(body, ensure_ascii=False).encode() if body is not None else None
    req = Request(BASE_URL + path, data=payload, headers={"Content-Type": "application/json"},
                  method="POST" if payload is not None else "GET")
    with urlopen(req, timeout=timeout) as r:
        return json.load(r)


def validate_rows(rows):
    """서버 smartstore_import.validate_reviews 와 동일한 규칙 (로컬 사전검증)."""
    if not rows:
        raise ValueError("후기 데이터가 비어 있습니다.")
    for i, row in enumerate(rows, 1):
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(row.get("date", ""))):
            raise ValueError(f"{i}번째 후기 날짜 형식 오류: {row.get('date')!r}")
        date.fromisoformat(row["date"])
        s = row.get("score")
        if isinstance(s, bool) or not isinstance(s, (int, float)) or not 1 <= s <= 5:
            raise ValueError(f"{i}번째 후기 평점 오류: {s!r}")
        for k in ("content", "author", "product"):
            if not isinstance(row.get(k, ""), str):
                raise ValueError(f"{i}번째 후기 {k} 형식 오류")
        if not row.get("content", "").strip():
            raise ValueError(f"{i}번째 후기 내용이 비어 있습니다")
        if row.get("platform") not in ("naver", "smartstore"):
            raise ValueError(f"{i}번째 후기 플랫폼 오류")
        row.pop("row_index", None)
    return rows


def upload(rows):
    """서버의 청크 업로드 → 병합 → 검증. 서버가 (review_no / author+date+content) 기준으로 중복 제거한다."""
    rows = validate_rows(rows)
    sid = uuid4().hex
    for off in range(0, len(rows), 500):
        r = api("/api/import-smartstore-chunk", dict(reviews=rows[off:off + 500], replace=off == 0, import_id=sid))
        if not r.get("ok") or r.get("total") != min(off + 500, len(rows)):
            raise RuntimeError("청크 전송 검증 실패 — 병합은 실행하지 않았습니다.")
    r = api(f"/api/import-smartstore-done?import_id={sid}&expected_count={len(rows)}", {})
    if not r.get("ok"):
        raise RuntimeError(f"서버 병합 실패: {r}")
    return r


def server_since_date(days_back=7, default_days=30):
    """서버의 마지막 후기 날짜 - days_back 부터 재수집 (늦게 달리는 후기·수정 반영)."""
    try:
        st = api("/api/smartstore-status")
        last = st.get("last_review_date")
        if last:
            return (date.fromisoformat(last) - timedelta(days=days_back)).isoformat()
    except Exception as e:
        log(f"서버 상태 조회 실패({e}) → 기본 {default_days}일 범위")
    return (date.today() - timedelta(days=default_days)).isoformat()


# ───────────────────────── 브라우저 ─────────────────────────
async def open_browser(pw, headless):
    PROFILE_DIR.mkdir(parents=True, exist_ok=True)
    DOWNLOAD_DIR.mkdir(parents=True, exist_ok=True)
    # persistent context = 크롬 프로필처럼 쿠키/로그인유지 토큰이 디스크에 남는다.
    return await pw.chromium.launch_persistent_context(
        str(PROFILE_DIR), headless=headless, locale="ko-KR", timezone_id="Asia/Seoul",
        viewport={"width": 1440, "height": 1000}, accept_downloads=True,
        args=["--disable-dev-shm-usage"],
    )


async def ensure_logged_in(page, wait_ms=30000):
    """판매자센터 좌측 메뉴(#seller-lnb)에 백년화편이 보이면 로그인 상태."""
    try:
        await page.locator("#seller-lnb").wait_for(timeout=wait_ms)
    except Exception:
        raise SessionExpired("판매자센터 메뉴가 뜨지 않음 (로그인 필요)")
    if "nid.naver.com" in page.url:
        raise SessionExpired("로그인 페이지로 이동됨")
    text = await page.locator("#seller-lnb").inner_text()
    if STORE_NAME not in text:
        raise SessionExpired(f"'{STORE_NAME}' 판매자 계정이 아님 또는 로그인 만료")


async def goto_review_page(page):
    await page.goto(REVIEW_URL, wait_until="domcontentloaded")
    await ensure_logged_in(page)
    if "/review/search" not in page.url:      # 첫 진입 시 홈으로 튕기는 경우
        await page.goto(REVIEW_URL, wait_until="domcontentloaded")
    await page.get_by_role("heading", name="리뷰관리", exact=True).wait_for(timeout=30000)
    close_btn = page.locator("#seller-rnb").get_by_role("button", name="닫기", exact=True)
    if await close_btn.count() and await close_btn.is_visible():
        await close_btn.click()


async def set_date_range(page, since):
    """'오늘' 버튼(양쪽 오늘로 세팅) → 시작일만 since 로 교체."""
    await page.get_by_role("button", name="오늘", exact=True).click()
    start = page.locator('input[title="날짜 입력"]').nth(0)
    await start.evaluate('(e)=>e.removeAttribute("readonly")')
    want = date.fromisoformat(since).strftime("%Y.%m.%d.")
    await start.fill(want)
    await start.press("Tab")
    if await start.input_value() != want:
        raise ValueError("시작일 설정 확인 실패")
    await page.get_by_role("button", name="검색", exact=True).click()
    await page.wait_for_timeout(4000)


async def read_total(page):
    heading = await page.get_by_role("heading", name=re.compile("리뷰목록")).inner_text()
    m = re.search(r"총\s*([\d,]+)\s*개", heading)
    if not m:
        raise ValueError("조회 건수 확인 실패")
    return int(m.group(1).replace(",", ""))


async def read_rows_ui(page, since):
    """① 리뷰관리 ag-grid 화면을 스크롤+페이지 넘김으로 전부 읽는다."""
    await goto_review_page(page)
    await set_date_range(page, since)
    expected = await read_total(page)
    if expected == 0:
        return []
    if expected > 10000:
        raise ValueError("조회 한도 초과(10,000건) — --since 로 범위를 줄이세요")
    found = {}
    for _page in range(500):
        vp = page.locator(".ag-body-viewport")
        await vp.scroll_into_view_if_needed()
        await vp.evaluate("(e)=>{e.scrollTop=0}")
        await page.wait_for_timeout(300)
        for _ in range(600):
            for row in await page.evaluate(ROW_JS):
                found[row["review_no"]] = row
            if len(found) >= expected:
                break
            at_end = await vp.evaluate("(e)=>e.scrollTop+e.clientHeight>=e.scrollHeight-2")
            if at_end:
                break
            await vp.evaluate("(e)=>{e.scrollTop+=Math.max(100,e.clientHeight*0.8)}")
            await page.wait_for_timeout(250)
        if len(found) >= expected:
            break
        nxt = page.locator('[aria-label="다음 페이지로 이동"]').filter(visible=True)
        if await nxt.count() == 0 or not await nxt.is_enabled():
            break
        await nxt.click()
        for _ in range(60):
            await page.wait_for_timeout(500)
            incoming = await page.evaluate(ROW_JS)
            if any(r["review_no"] not in found for r in incoming):
                break
        else:
            raise ValueError("다음 페이지에 새 후기가 표시되지 않음")
    # 헤더 총건수가 이전 검색값을 물고 있는 경우가 있어 최종값으로 재대조
    final_total = await read_total(page)
    if len(found) != final_total and len(found) != expected:
        raise ValueError(f"조회 건수({final_total})와 수집 건수({len(found)}) 불일치")
    rows = [r for r in found.values() if r["date"] >= since]
    return validate_rows(rows)


async def read_rows_excel(page, since):
    """② 폴백: 같은 화면에서 '엑셀 다운로드' → openpyxl 파싱. 버튼/팝업 문구는 사이트 변경에 따라 조정 필요."""
    from openpyxl import load_workbook
    await goto_review_page(page)
    await set_date_range(page, since)
    btn = page.get_by_role("button", name=re.compile(r"엑셀\s*다운")).filter(visible=True)
    if await btn.count() == 0:
        raise ValueError("엑셀 다운로드 버튼을 찾지 못함")
    async with page.expect_download(timeout=120000) as dl_info:
        await btn.first.click()
        # 다운로드 사유/확인 팝업이 뜨는 경우 (있으면 눌러주고, 없으면 무시)
        for name in ("확인", "다운로드"):
            ok = page.get_by_role("button", name=name, exact=True).filter(visible=True)
            try:
                if await ok.count():
                    await ok.first.click(timeout=3000)
            except Exception:
                pass
    dl = await dl_info.value
    path = DOWNLOAD_DIR / f"reviews_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    await dl.save_as(str(path))
    log(f"엑셀 저장: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        ws.reset_dimensions()
        it = iter(ws.iter_rows(values_only=True))
        headers = next(it)
        idx = {str(v).strip(): i for i, v in enumerate(headers) if v is not None}
        need = ["리뷰등록일", "구매자평점", "상품명", "리뷰상세내용", "등록자", "리뷰글번호"]
        if not all(k in idx for k in need):
            raise ValueError(f"엑셀 열 구성이 예상과 다름: {list(idx)}")
        out = []
        for row in it:
            if not any(v is not None for v in row):
                continue
            c = lambda k: row[idx[k]] if k in idx and idx[k] < len(row) else None
            raw = c("리뷰등록일")
            if isinstance(raw, datetime):
                d = raw.strftime("%Y-%m-%d")
            else:
                m = re.match(r"\s*(\d{4})[.\-/](\d{2})[.\-/](\d{2})", str(raw))
                if not m:
                    raise ValueError(f"날짜 형식 확인 불가: {raw!r}")
                d = "-".join(m.groups())
            if d < since:
                continue
            out.append(dict(date=d, score=float(c("구매자평점") or 0), product=str(c("상품명") or ""),
                            content=str(c("리뷰상세내용") or ""), author=str(c("등록자") or ""),
                            review_no=str(c("리뷰글번호") or ""), review_type=str(c("리뷰구분") or ""),
                            platform="naver", title=""))
        return validate_rows(out) if out else []
    finally:
        wb.close()


# ───────────────────────── 실행 모드 ─────────────────────────
def push_session(storage, token):
    """로그인된 브라우저 세션(storage_state)을 서버로 이전. 서버가 검증 후 저장하고 첫 수집을 시작한다."""
    payload = json.dumps(storage).encode()
    req = Request(BASE_URL + "/api/naver-seller/session", data=payload,
                  headers={"Content-Type": "application/json", "X-Naver-Connect": token}, method="POST")
    try:
        with urlopen(req, timeout=180) as r:
            return json.load(r)
    except HTTPError as e:
        raise RuntimeError(f"서버 응답 {e.code}: {e.read().decode('utf-8', 'ignore')[:300]}")


async def do_login(push=False):
    from playwright.async_api import async_playwright
    token = None
    if push:
        token = os.environ.get("NAVER_CONNECT_TOKEN") or input("서버 연결 토큰(NAVER_CONNECT_TOKEN): ").strip()
        if not token:
            print("토큰이 없어 --push 를 건너뜁니다."); push = False
    async with async_playwright() as pw:
        ctx = await open_browser(pw, headless=False)
        page = ctx.pages[0] if ctx.pages else await ctx.new_page()
        await page.goto(SELLER_HOME, wait_until="domcontentloaded")
        print("\n▶ 브라우저 창에서 판매자센터에 직접 로그인하세요. ('로그인 상태 유지' 체크 권장)")
        print("  로그인이 확인되면 이 창은 자동으로 닫힙니다. (최대 10분 대기)\n")
        deadline = time.time() + 600
        while time.time() < deadline:
            try:
                await ensure_logged_in(page, wait_ms=3000)
                log("로그인 확인 — 세션이 프로필에 저장되었습니다.")
                await page.wait_for_timeout(1500)
                if push:
                    storage = await ctx.storage_state()
                    log("서버로 세션 이전 중 (서버가 검증까지 하므로 1분 정도 걸릴 수 있음)...")
                    r = push_session(storage, token)
                    log(f"서버 응답: {r.get('message', r)}")
                await ctx.close()
                try:
                    api("/api/smartstore-cookie-ok", {})
                except Exception:
                    pass
                return 0
            except SessionExpired:
                await page.wait_for_timeout(2000)
        log("10분 안에 로그인이 확인되지 않아 종료합니다.")
        await ctx.close()
        return 2


async def do_collect(args):
    from playwright.async_api import async_playwright
    since = args.since or server_since_date(days_back=args.overlap_days)
    log(f"수집 시작 · since={since} · 서버={BASE_URL}")
    result = {"since": since, "started": datetime.now().isoformat()}
    async with async_playwright() as pw:
        ctx = await open_browser(pw, headless=args.headless)
        page = ctx.pages[0] if ctx.pages else await ctx.new_page()
        page.set_default_timeout(20000)
        try:
            rows, source = None, None
            try:
                rows = await read_rows_ui(page, since)
                source = "seller_ui"
            except SessionExpired:
                raise
            except Exception as e:
                log(f"화면 읽기 실패({type(e).__name__}: {e}) → 엑셀 폴백 시도")
                rows = await read_rows_excel(page, since)
                source = "seller_excel"
            result.update(source=source, received=len(rows))
            if not rows:
                log("해당 기간 후기 0건 — 업로드 생략")
                result.update(added=0, ok=True)
            else:
                r = upload(rows)
                result.update(added=r.get("added"), updated=r.get("updated"), ok=True)
                log(f"업로드 완료 · 수신 {len(rows)}건 · 신규 {r.get('added')}건 · 갱신 {r.get('updated')}건 ({source})")
            try:
                api("/api/smartstore-cookie-ok", {})
            except Exception:
                pass
            return 0
        except SessionExpired as e:
            log(f"세션 만료: {e} → 'python naver_seller_local.py --login' 을 다시 실행하세요")
            result.update(ok=False, error="session_expired")
            try:
                api("/api/smartstore-cookie-expired", {"expired_at": datetime.now().isoformat()})
            except Exception:
                pass
            return 2
        except Exception as e:
            log(f"수집 실패: {type(e).__name__}: {e}")
            result.update(ok=False, error=f"{type(e).__name__}: {e}")
            return 1
        finally:
            result["finished"] = datetime.now().isoformat()
            try:
                LAST_RESULT.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
            except OSError:
                pass
            await ctx.close()


def main():
    p = argparse.ArgumentParser(description="백년화편 네이버 판매자센터 후기 수집기 (회사 PC용)")
    p.add_argument("--login", action="store_true", help="브라우저를 띄워 수동 로그인 후 세션 저장")
    p.add_argument("--push", action="store_true", help="--login 과 함께: 로그인 세션을 서버로 이전 (이후 서버가 매일 수집)")
    p.add_argument("--since", help="수집 시작일 YYYY-MM-DD (기본: 서버 마지막 후기일 - overlap)")
    p.add_argument("--overlap-days", type=int, default=7, help="서버 마지막 후기일에서 며칠 전부터 다시 읽을지 (기본 7)")
    p.add_argument("--headless", action="store_true", help="창 없이 실행 (세션 안정성은 창 있는 쪽이 더 낫다)")
    a = p.parse_args()
    if sys.platform == "win32":
        asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())
    code = asyncio.run(do_login(push=a.push) if a.login else do_collect(a))
    sys.exit(code)


if __name__ == "__main__":
    main()
