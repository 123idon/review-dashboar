"""Deterministic SmartStore export worker. No model/API calls to an LLM."""
import asyncio
from datetime import datetime
import hashlib
import hmac
import json
import os
from pathlib import Path
import re
import tempfile
import time
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

from smartstore_import import validate_reviews

KST = ZoneInfo('Asia/Seoul')
REVIEW_URL = 'https://sell.smartstore.naver.com/#/review/search'
DATA = Path(os.environ.get('NAVER_DATA_DIR', 'data'))
AUTH = DATA / 'naver_private' / 'state.json'
STATUS = DATA / 'naver_automation_status.json'
LOCK = asyncio.Lock()


def now():
    return datetime.now(KST).isoformat()


def private_save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    fd, name = tempfile.mkstemp(dir=path.parent)
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as stream:
            json.dump(value, stream, ensure_ascii=False)
        os.replace(name, path)
        path.chmod(0o600)
    finally:
        Path(name).unlink(missing_ok=True)


def status():
    try:
        result = json.loads(STATUS.read_text(encoding='utf-8'))
    except (OSError, ValueError):
        result = {'state': 'needs_login', 'message': '최초 네이버 로그인 연결이 필요합니다.'}
    result.update(schedule='매일 00:00', timezone='Asia/Seoul', gpt_tokens=0,
                  session_configured=AUTH.exists(), running=LOCK.locked())
    if not AUTH.exists() and not LOCK.locked():
        result['state'] = 'needs_login'
    elif result.get('state') == 'running' and not LOCK.locked():
        result.update(state='interrupted', message='이전 수집이 서버 재시작으로 중단되었습니다.')
    return result


def record(state, message, **extra):
    old = status()
    old.update(state=state, message=message, checked_at=now(), **extra)
    private_save(STATUS, old)


def pairing_allowed(token):
    expected = os.environ.get('NAVER_PAIRING_SHA256', '')
    if not expected or time.time() > float(os.environ.get('NAVER_PAIRING_EXPIRES', '0')):
        return False
    used = DATA / 'naver_private' / 'paired.json'
    if used.exists() and json.loads(used.read_text()).get('hash') == expected:
        return False
    return hmac.compare_digest(hashlib.sha256(token.encode()).hexdigest(), expected)


def filter_state(state):
    """Accept only Naver session state; never persist unrelated site credentials."""
    if not isinstance(state, dict):
        raise ValueError('잘못된 로그인 자료입니다.')
    def naver(host):
        host = str(host).lstrip('.').lower()
        return host == 'naver.com' or host.endswith('.naver.com')
    cookies = [c for c in state.get('cookies', []) if isinstance(c, dict) and naver(c.get('domain', ''))]
    origins = [o for o in state.get('origins', []) if isinstance(o, dict)
               and urlparse(o.get('origin', '')).scheme == 'https'
               and naver(urlparse(o.get('origin', '')).hostname or '')]
    if not cookies:
        raise ValueError('네이버 로그인 상태가 없습니다.')
    return {'cookies': cookies, 'origins': origins}


def parse_export(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        sheet.reset_dimensions()
        rows = iter(sheet.iter_rows(values_only=True))
        headers = next(rows)
        index = {str(v).strip(): i for i, v in enumerate(headers) if v is not None}
        required = ['리뷰등록일', '구매자평점', '상품명', '리뷰상세내용', '등록자', '리뷰글번호']
        if not all(k in index for k in required):
            raise ValueError('판매자센터 리뷰 엑셀 열이 변경되었습니다.')
        out = []
        for row in rows:
            if not any(v is not None for v in row):
                continue
            def cell(key):
                i = index.get(key)
                return row[i] if i is not None and i < len(row) else None
            raw_date = cell('리뷰등록일')
            if isinstance(raw_date, datetime):
                date = raw_date.strftime('%Y-%m-%d')
            else:
                match = re.match(r'\s*(\d{4})[.\-/](\d{2})[.\-/](\d{2})', str(raw_date))
                if not match:
                    raise ValueError('리뷰 날짜 형식을 확인할 수 없습니다.')
                date = '-'.join(match.groups())
            out.append(dict(date=date, score=cell('구매자평점'), product=str(cell('상품명') or ''),
                            content=str(cell('리뷰상세내용') or ''), author=str(cell('등록자') or ''),
                            review_no=str(cell('리뷰글번호') or ''), review_type=str(cell('리뷰구분') or ''),
                            platform='naver', title=''))
        return validate_reviews(out)
    finally:
        workbook.close()


async def export_reviews(state):
    """Use the seller's visible review export; never bypass login or bot checks."""
    from playwright.async_api import async_playwright
    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=True, args=['--disable-dev-shm-usage'])
        try:
            context = await browser.new_context(storage_state=state, accept_downloads=True,
                                                locale='ko-KR', timezone_id='Asia/Seoul')
            page = await context.new_page()
            page.set_default_timeout(30000)
            page.on('dialog', lambda dialog: dialog.accept() if dialog.type in ('confirm', 'alert') else dialog.dismiss())
            await page.goto(REVIEW_URL, wait_until='domcontentloaded', timeout=60000)
            # These labels were verified on the seller's review-management page.
            await page.get_by_text('6개월', exact=True).wait_for(timeout=60000)
            if '백년화편' not in await page.locator('body').inner_text():
                raise ValueError('백년화편 스토어 화면인지 확인하지 못했습니다.')
            await page.get_by_text('6개월', exact=True).click()
            await page.get_by_role('button', name='검색', exact=True).click()
            await page.get_by_text('엑셀다운', exact=True).wait_for()
            async with page.expect_download(timeout=180000) as pending:
                await page.get_by_text('엑셀다운', exact=True).click()
                # Some versions present an HTML confirmation instead of a JS dialog.
                try:
                    confirm = page.get_by_role('button', name='확인', exact=True)
                    await confirm.wait_for(state='visible', timeout=3000)
                    await confirm.click()
                except Exception:
                    pass
            download = await pending.value
            with tempfile.TemporaryDirectory(prefix='naver-export-') as directory:
                path = Path(directory) / 'reviews.xlsx'
                await download.save_as(path)
                reviews = await asyncio.to_thread(parse_export, path)
            refreshed = filter_state(await context.storage_state())
            return reviews, refreshed
        finally:
            await browser.close()
